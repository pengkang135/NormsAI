# -*- coding: utf-8 -*-
"""价格库 → Excel 导出。

按当前所在页面导出对应内容，三种形态：

- `list`     资源价格清单（分册 / 配比材料 / 机械台班 / 机电主材 / 全部），带完整溯源列
- `coverage` 覆盖报告（册×类别矩阵 + 来源可靠度 + 未定价缺口）
- `gap`      缺口询价表（未定价条目 + 候选报价），可直接拿去询价

每份都附一张「导出说明」记录价格包口径与筛选条件——脱离前端后仍能说清这批数是怎么来的。
样式沿用 src/export_norms.py 的字体/配色约定。
"""
import json
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from src import price_api as PA

FONT_NAME = '微软雅黑'
FONT_HDR = Font(name=FONT_NAME, size=9, bold=True)
FONT_DATA = Font(name=FONT_NAME, size=9)
FONT_GRAY = Font(name=FONT_NAME, size=9, color='808080')
FONT_TITLE = Font(name=FONT_NAME, size=11, bold=True)
FONT_SUB = Font(name=FONT_NAME, size=9, bold=True, color='1A1A1A')

FILL_HDR = PatternFill('solid', fgColor='FFD9E1F2')
FILL_SUB = PatternFill('solid', fgColor='FFFBE5D6')
FILL_BAD = PatternFill('solid', fgColor='FFF8D7DA')     # 覆盖率 <40%
FILL_WARN = PatternFill('solid', fgColor='FFFFF3CD')    # 覆盖率 40~80%
FILL_OK = PatternFill('solid', fgColor='FFE2EFDA')      # 覆盖率 >=80%

THIN = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))

NUM = '#,##0.00'
INT = '#,##0'

SCOPE_LABEL = {
    'vol': '分册', 'mix': '配比材料', 'machine': '机械台班',
    'me_main': '机电主材', 'all': '全部人材机',
}

LIST_HEADERS = [
    '序号', '编码', '类别', '名称', '规格', '单位',
    '名称(EN)', '规格(EN)', '单位(EN)',
    '含税单价', '不含税单价', '币种',
    '来源等级', '来源类型', '定价日期',
    '原始报价', '原币种', '原单位', '汇率', '单位换算',
    '供应商', '来源文件/项目', '推导依据',
    '定额引用次数', '所属册',
]
LIST_WIDTHS = [6, 13, 7, 26, 18, 8, 26, 18, 10, 12, 12, 7, 9, 10, 12,
               12, 9, 9, 8, 10, 16, 34, 50, 11, 9]

GAP_HEADERS = [
    '序号', '缺口编码', '类别', '缺口名称', '规格', '单位', '名称(EN)',
    '定额引用次数', '所属册',
    '候选名称', '候选规格', '候选单位', '单位是否一致', '相似度',
    '候选含税价', '币种', '报价日期', '供应商', '项目', '来源文件', '专业', '来源等级',
]
GAP_WIDTHS = [6, 13, 7, 24, 16, 8, 22, 11, 9,
              26, 20, 10, 12, 8, 12, 8, 12, 16, 22, 32, 12, 9]


def _hdr_row(ws, headers, widths, row=1):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row, ci, value=h)
        c.font = FONT_HDR
        c.fill = FILL_HDR
        c.border = THIN
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1).coordinate


def _style_row(ws, row, ncols, font=FONT_DATA):
    for ci in range(1, ncols + 1):
        c = ws.cell(row, ci)
        c.font = font
        c.border = THIN


def _cov_fill(ratio):
    if ratio >= 0.8:
        return FILL_OK
    return FILL_WARN if ratio >= 0.4 else FILL_BAD


def _clean_country(v):
    """部分 CostSpread 行的 country 损坏成字面 '??'，不写假国别。"""
    return '' if (not v or v == '??') else v


def _parse_project(features):
    """43 条直采报价把项目/地点塞在 features 里：project=… | location=… | sheet=…"""
    if not features or 'project=' not in features:
        return '', features or ''
    project = location = ''
    rest = []
    for seg in features.split('|'):
        seg = seg.strip()
        if seg.startswith('project='):
            project = seg[len('project='):].strip()
        elif seg.startswith('location='):
            location = seg[len('location='):].strip()
        elif seg and '=' not in seg:
            rest.append(seg)
    return (project + (' · ' + location if location else '')), ' '.join(rest)


# ── 资源价格清单 ──────────────────────────────────────────────────────────

LIST_SQL = """
SELECT u.res_ID, r.res_Code, r.res_Class_Name,
       r.res_Name, r.res_Standard, r.res_Unit,
       r.res_Name_EN, r.res_Standard_EN, r.res_Unit_EN,
       p.price, p.price_excl, p.currency, p.tier, p.price_kind, p.price_date,
       p.src_price, p.src_currency, p.src_unit, p.fx_rate, p.unit_factor,
       p.supplier, p.source_doc, p.source,
       SUM(u.refs) AS refs,
       GROUP_CONCAT(DISTINCT u.vol) AS vols,
       MIN(u.cons_Style) AS cons_Style
{base}
GROUP BY u.res_ID
ORDER BY (p.price IS NULL) DESC, refs DESC, r.res_Code
"""


def _sheet_list(wb, conn, pack, params, fx_cny):
    base, args = PA.build_list_query(
        pack, params.get('scope', 'all'), params.get('vol'), params.get('cls'),
        params.get('priced'), params.get('tier'), params.get('q'))
    rows = [dict(r) for r in conn.execute(LIST_SQL.format(base=base), args)]

    ws = wb.create_sheet('人材机价格清单')
    _hdr_row(ws, LIST_HEADERS, LIST_WIDTHS)

    priced = 0
    for i, r in enumerate(rows):
        rr = i + 2
        has = r['price'] is not None
        if has:
            priced += 1
        cls = PA.CLS_NAME.get(r['cons_Style'], (r['res_Class_Name'] or '', ''))[0]
        ws.cell(rr, 1, value=i + 1).alignment = Alignment(horizontal='center')
        ws.cell(rr, 2, value=r['res_Code'])
        ws.cell(rr, 3, value=cls)
        ws.cell(rr, 4, value=r['res_Name'])
        ws.cell(rr, 5, value=r['res_Standard'])
        ws.cell(rr, 6, value=r['res_Unit'])
        ws.cell(rr, 7, value=r['res_Name_EN'])
        ws.cell(rr, 8, value=r['res_Standard_EN'])
        ws.cell(rr, 9, value=r['res_Unit_EN'])
        if has:
            ws.cell(rr, 10, value=round(r['price'], 4)).number_format = NUM
            if r['price_excl'] is not None:
                ws.cell(rr, 11, value=round(r['price_excl'], 4)).number_format = NUM
            ws.cell(rr, 12, value=r['currency'])
            ws.cell(rr, 13, value=r['tier'])
            ws.cell(rr, 14, value=r['price_kind'])
            ws.cell(rr, 15, value=r['price_date'])
            if r['src_price'] is not None:
                ws.cell(rr, 16, value=round(r['src_price'], 4)).number_format = NUM
            ws.cell(rr, 17, value=r['src_currency'])
            ws.cell(rr, 18, value=r['src_unit'])
            if r['fx_rate'] is not None:
                ws.cell(rr, 19, value=round(r['fx_rate'], 6))
            if r['unit_factor'] is not None:
                ws.cell(rr, 20, value=round(r['unit_factor'], 6))
            ws.cell(rr, 21, value=r['supplier'])
            ws.cell(rr, 22, value=r['source_doc'])
            src = ws.cell(rr, 23, value=r['source'])
            src.alignment = Alignment(wrap_text=True, vertical='top')
        else:
            ws.cell(rr, 13, value='待询价')
        ws.cell(rr, 24, value=r['refs']).number_format = INT
        ws.cell(rr, 25, value=''.join(sorted((r['vols'] or '').split(','))))
        _style_row(ws, rr, len(LIST_HEADERS), FONT_DATA if has else FONT_GRAY)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(LIST_HEADERS))}{max(len(rows) + 1, 2)}"
    return {'rows': len(rows), 'priced': priced}


# ── 覆盖报告 ──────────────────────────────────────────────────────────────

def _sheet_coverage(wb, conn, pack, fx_cny, gap_limit=500):
    data, status, err = PA.api_coverage(pack, gap_limit=gap_limit)
    if err:
        raise RuntimeError(err)

    ws = wb.create_sheet('覆盖矩阵')
    s = data['summary']
    ws.cell(1, 1, value='册 × 类别 价格覆盖率').font = FONT_TITLE
    ws.cell(2, 1, value=(
        f"整体 {s['priced']}/{s['total']} 条目有价 "
        f"({s['priced'] / s['total'] * 100:.1f}%)，"
        f"按定额引用次数加权 {s['refs_priced'] / s['refs'] * 100:.1f}% "
        f"({s['refs_priced']:,}/{s['refs']:,})")).font = FONT_GRAY

    clss = sorted({m['cls'] for m in data['matrix']})
    vols = []
    cell = {}
    for m in data['matrix']:
        if m['vol'] not in vols:
            vols.append(m['vol'])
        cell[(m['vol'], m['cls'])] = m

    top = 4
    ws.cell(top, 1, value='册').font = FONT_HDR
    ws.cell(top, 1).fill = FILL_HDR
    ws.cell(top, 1).border = THIN
    for ci, c in enumerate(clss, 2):
        h = ws.cell(top, ci, value=PA.CLS_NAME.get(c, ('', ''))[0])
        h.font = FONT_HDR
        h.fill = FILL_HDR
        h.border = THIN
        h.alignment = Alignment(horizontal='center')
    ws.cell(top, len(clss) + 2, value='合计').font = FONT_HDR
    ws.cell(top, len(clss) + 2).fill = FILL_HDR
    ws.cell(top, len(clss) + 2).border = THIN

    for ri, v in enumerate(vols, top + 1):
        name = PA.VOL_NAME.get(v, (v, ''))[0]
        c = ws.cell(ri, 1, value=name)
        c.font = FONT_SUB
        c.fill = FILL_SUB
        c.border = THIN
        vt = vp = 0
        for ci, cl in enumerate(clss, 2):
            m = cell.get((v, cl))
            cc = ws.cell(ri, ci)
            cc.border = THIN
            cc.font = FONT_DATA
            cc.alignment = Alignment(horizontal='center')
            if not m:
                cc.value = '—'
                continue
            vt += m['total']
            vp += m['priced']
            cc.value = f"{m['priced']}/{m['total']}  {m['priced'] / m['total'] * 100:.1f}%"
            cc.fill = _cov_fill(m['priced'] / m['total'])
        tc = ws.cell(ri, len(clss) + 2,
                     value=f"{vp}/{vt}  {vp / vt * 100:.1f}%" if vt else '—')
        tc.font = FONT_SUB
        tc.border = THIN
        tc.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 18
    for ci in range(2, len(clss) + 3):
        ws.column_dimensions[get_column_letter(ci)].width = 17

    trow = top + len(vols) + 3
    ws.cell(trow, 1, value='来源可靠度分布').font = FONT_TITLE
    trow += 1
    for ci, h in enumerate(['等级', '含义', '定价方式', '条数'], 1):
        c = ws.cell(trow, ci, value=h)
        c.font = FONT_HDR
        c.fill = FILL_HDR
        c.border = THIN
    for tier in ['T1', 'T2', 'T3', 'T4', 'S', 'D']:
        rows = [x for x in data['tiers'] if x['tier'] == tier]
        if not rows:
            continue
        trow += 1
        ws.cell(trow, 1, value=tier)
        ws.cell(trow, 2, value=PA.TIER_NAME.get(tier, ('', ''))[0])
        ws.cell(trow, 3, value=', '.join(sorted({x['price_kind'] or '' for x in rows})))
        ws.cell(trow, 4, value=sum(x['n'] for x in rows)).number_format = INT
        _style_row(ws, trow, 4)

    _sheet_gap_list(wb, data['gaps'])
    return {'vols': len(vols), 'gaps': len(data['gaps'])}


def _sheet_gap_list(wb, gaps):
    ws = wb.create_sheet('未定价缺口')
    headers = ['序号', '编码', '类别', '名称', '规格', '单位', '名称(EN)', '单位(EN)',
               '定额引用次数', '所属册']
    _hdr_row(ws, headers, [6, 13, 7, 28, 20, 8, 26, 10, 12, 9])
    for i, g in enumerate(gaps):
        rr = i + 2
        ws.cell(rr, 1, value=i + 1).alignment = Alignment(horizontal='center')
        ws.cell(rr, 2, value=g['res_Code'])
        ws.cell(rr, 3, value=g['res_Class_Name'])
        ws.cell(rr, 4, value=g['res_Name'])
        ws.cell(rr, 5, value=g['res_Standard'])
        ws.cell(rr, 6, value=g['res_Unit'])
        ws.cell(rr, 7, value=g['res_Name_EN'])
        ws.cell(rr, 8, value=g['res_Unit_EN'])
        ws.cell(rr, 9, value=g['refs']).number_format = INT
        ws.cell(rr, 10, value=''.join(g.get('vols') or []))
        _style_row(ws, rr, len(headers))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(gaps) + 1, 2)}"


# ── 缺口询价表 ────────────────────────────────────────────────────────────

def _sheet_gap(wb, pack, params):
    data, status, err = PA.api_gap(
        pack=pack, vol=params.get('vol'), cls=params.get('cls'),
        scope=params.get('scope', 'me_main'),
        limit=int(params.get('gap_limit') or 200), cand=int(params.get('cand') or 5))
    if err:
        raise RuntimeError(err)

    ws = wb.create_sheet('缺口询价表')
    _hdr_row(ws, GAP_HEADERS, GAP_WIDTHS)

    rr = 1
    for i, it in enumerate(data['items']):
        cands = it['candidates'] or [None]
        first = rr + 1
        for c in cands:
            rr += 1
            ws.cell(rr, 1, value=i + 1).alignment = Alignment(horizontal='center')
            ws.cell(rr, 2, value=it['res_Code'])
            ws.cell(rr, 3, value=it['res_Class_Name'])
            ws.cell(rr, 4, value=it['res_Name'])
            ws.cell(rr, 5, value=it['res_Standard'])
            ws.cell(rr, 6, value=it['res_Unit'])
            ws.cell(rr, 7, value=it['res_Name_EN'])
            ws.cell(rr, 8, value=it['refs']).number_format = INT
            ws.cell(rr, 9, value=''.join(it.get('vols') or []))
            if c is None:
                n = ws.cell(rr, 10, value='全库无对应报价，需外部询价')
                n.font = Font(name=FONT_NAME, size=9, color='B8863A')
                _style_row(ws, rr, 9)
                for ci in range(10, len(GAP_HEADERS) + 1):
                    ws.cell(rr, ci).border = THIN
                continue
            project, spec = _parse_project(c.get('features'))
            ws.cell(rr, 10, value=(c.get('name') or '').replace('\n', ' ').strip())
            ws.cell(rr, 11, value=spec.replace('\n', ' ').strip()[:200])
            ws.cell(rr, 12, value=c.get('unit'))
            uc = ws.cell(rr, 13, value='一致' if c.get('unit_match') else '不一致')
            if not c.get('unit_match'):
                uc.font = Font(name=FONT_NAME, size=9, color='B8863A')
                uc.fill = FILL_WARN
            ws.cell(rr, 14, value=c.get('score'))
            if c.get('price_incl') is not None:
                ws.cell(rr, 15, value=round(c['price_incl'], 4)).number_format = NUM
            ws.cell(rr, 16, value=c.get('currency'))
            ws.cell(rr, 17, value=c.get('quote_date'))
            ws.cell(rr, 18, value=c.get('supplier'))
            ws.cell(rr, 19, value=project)
            ws.cell(rr, 20, value=c.get('source_doc'))
            ws.cell(rr, 21, value=c.get('specialty'))
            ws.cell(rr, 22, value=c.get('tier'))
            _style_row(ws, rr, len(GAP_HEADERS))
        # 同一缺口的多条候选：左侧缺口信息合并，右侧候选逐行排开
        if rr > first:
            for ci in range(1, 10):
                ws.merge_cells(start_row=first, start_column=ci, end_row=rr, end_column=ci)
                ws.cell(first, ci).alignment = Alignment(vertical='center',
                                                         horizontal='center' if ci in (1, 8, 9) else 'left')
    ws.auto_filter.ref = f"A1:{get_column_letter(len(GAP_HEADERS))}{max(rr, 2)}"
    stats = {'gaps': data['count'], 'matched': data['matched'], 'pool': data['pool_size']}
    if data.get('truncated'):
        stats['truncated'] = data['truncated']
        stats['gap_total'] = data['total']
    return stats


# ── 导出说明 ──────────────────────────────────────────────────────────────

def _sheet_meta(wb, pack_row, params, stats, view):
    ws = wb.create_sheet('导出说明', 0)
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 96

    ws.cell(1, 1, value='价格库导出').font = FONT_TITLE
    row = 3

    def kv(k, v):
        nonlocal row
        if v in (None, ''):
            return
        a = ws.cell(row, 1, value=k)
        a.font = FONT_HDR
        a.fill = FILL_HDR
        a.border = THIN
        b = ws.cell(row, 2, value=str(v))
        b.font = FONT_DATA
        b.border = THIN
        b.alignment = Alignment(wrap_text=True, vertical='top')
        row += 1

    kv('导出内容', {'list': '资源价格清单', 'coverage': '覆盖报告',
                    'gap': '缺口询价表'}.get(view, view))
    if pack_row:
        fx = ''
        try:
            fx = ', '.join(f'1 {pack_row["currency"]} → {1 / v:.4f} {k}'
                           for k, v in json.loads(pack_row['fx_json'] or '{}').items() if v)
        except (ValueError, TypeError, ZeroDivisionError):
            fx = pack_row['fx_json'] or ''
        kv('价格包', f"{pack_row['pack']} · {pack_row['name']}")
        kv('国别 / 币种', f"{pack_row['country'] or ''} / {pack_row['currency'] or ''}")
        kv('汇率', fx)
        kv('截止日期', pack_row['cutoff_date'])
        kv('状态', pack_row['status'])
        kv('口径说明', pack_row['note'])

    scope = params.get('scope', 'all')
    parts = [SCOPE_LABEL.get(scope, scope)]
    if params.get('vol'):
        parts.append(PA.VOL_NAME.get(params['vol'], (params['vol'], ''))[0])
    if params.get('cls'):
        parts.append(PA.CLS_NAME.get(int(params['cls']), ('', ''))[0])
    kv('范围', ' · '.join(p for p in parts if p))
    kv('有价筛选', {'yes': '仅有价', 'no': '仅无价'}.get(params.get('priced'), '全部'))
    kv('来源等级筛选', params.get('tier') or '全部')
    kv('关键字', params.get('q'))

    for k, v in (stats or {}).items():
        if k in ('truncated', 'gap_total'):
            continue
        kv({'rows': '导出条数', 'priced': '其中有价', 'gaps': '缺口条数',
            'matched': '有候选的缺口', 'pool': '候选报价池', 'vols': '册数'}.get(k, k),
           f'{v:,}' if isinstance(v, int) else v)
    if (stats or {}).get('truncated'):
        kv('⚠ 已截断', f"本范围共 {stats['gap_total']:,} 条缺口，"
                       f"按定额引用次数取前 {stats['gaps']:,} 条导出，"
                       f"另有 {stats['truncated']:,} 条未包含。"
                       f"撮合逐条比对 13,421 条报价，全量导出耗时过长，"
                       f"如需其余部分请缩小范围（按册/按类别）后再导出。")

    row += 1
    note = ws.cell(row, 1, value='说明')
    note.font = FONT_HDR
    note.fill = FILL_HDR
    note.border = THIN
    txt = ws.cell(row, 2, value=(
        '来源等级 T1 泰国直采 > T2 邻国报价 > T3 中国报价 > T4 非洲报价 > S 近似重名 > D 派生。'
        'D 为中国基线经汇率与标定系数派生，最粗糙，用于参考不用于结算。\n'
        '「待询价」表示该资源在本价格包中尚未定价，不是价格为 0。\n'
        '缺口询价表的候选由名称相似度排序，仅作询价线索；单位不一致的行需人工判断。'))
    txt.font = FONT_GRAY
    txt.border = THIN
    txt.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[row].height = 58


# ── 入口 ──────────────────────────────────────────────────────────────────

def build_price_workbook(params):
    """params: {view, pack, scope, vol, cls, priced, tier, q, gap_limit, cand}

    返回 (workbook, stats)。view 取 list / coverage / gap。
    """
    view = params.get('view') or 'list'
    conn = PA._conn()
    if not conn:
        raise RuntimeError('resource_master.sqlite 不存在')
    try:
        pack = params.get('pack') or PA._default_pack(conn)
        scope = params.get('scope', 'all')
        if scope not in PA.SCOPE_WHERE:
            raise RuntimeError(f'未知 scope: {scope}')
        pack_row = conn.execute(
            'SELECT * FROM price_pack WHERE pack = ?', (pack,)).fetchone()
        fx_cny = 5.0
        if pack_row:
            try:
                fx_cny = json.loads(pack_row['fx_json'] or '{}').get('CNY') or 5.0
            except (ValueError, TypeError):
                pass

        wb = Workbook()
        wb.remove(wb.active)

        if view == 'coverage':
            stats = _sheet_coverage(wb, conn, pack, fx_cny)
        elif view == 'gap':
            stats = _sheet_gap(wb, pack, params)
        else:
            view = 'list'
            stats = _sheet_list(wb, conn, pack, params, fx_cny)

        _sheet_meta(wb, pack_row, params, stats, view)
        return wb, stats
    finally:
        conn.close()


def export_filename(params):
    view = params.get('view') or 'list'
    if view == 'coverage':
        return '价格库_覆盖报告.xlsx'
    scope = params.get('scope', 'all')
    parts = []
    if params.get('vol'):
        parts.append(PA.VOL_NAME.get(params['vol'], (params['vol'], ''))[0])
        if params.get('cls'):
            parts.append(PA.CLS_NAME.get(int(params['cls']), ('', ''))[0])
    else:
        parts.append(SCOPE_LABEL.get(scope, scope))
    tag = '缺口询价表' if view == 'gap' else '价格清单'
    # 筛选条件进文件名，否则同一页面筛选前后导出会互相覆盖
    if view == 'list':
        if params.get('priced') == 'yes':
            parts.append('仅有价')
        elif params.get('priced') == 'no':
            parts.append('仅无价')
        if params.get('tier'):
            parts.append(params['tier'].replace(',', ''))
        if params.get('q'):
            parts.append(str(params['q'])[:12])
    return f"价格库_{'_'.join(p for p in parts if p)}_{tag}.xlsx"
